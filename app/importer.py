"""CSV / Excel -> vulnerabilities table loader."""
import csv
import io
import re
from datetime import date, datetime

from openpyxl import load_workbook

from dateutil import parser as dateparser
from sqlalchemy import delete, select, text
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.exc import IntegrityError

from .models import VULN_FIELDS, ImportRun, StagedVulnerability, UploadLog

BATCH_SIZE = 1000


def _norm(header):
    return re.sub(r"[^a-z0-9]", "", (header or "").lower())


# Normalised header -> (column, kind). Also accepts the "Enviroment" spelling.
HEADER_MAP = {_norm(h): (name, kind) for h, name, kind in VULN_FIELDS}
HEADER_MAP["enviroment"] = ("environment", "str")


EXCEL_EXTENSIONS = (".xlsx", ".xlsm")
SUPPORTED_EXTENSIONS = (".csv",) + EXCEL_EXTENSIONS


def _convert(value, kind):
    # Excel cells arrive already typed (datetime, int, float); CSV cells are strings.
    if isinstance(value, datetime):
        return value.date() if kind == "date" else value.replace(tzinfo=None) if kind == "datetime" else value.isoformat(" ")
    if isinstance(value, date):
        return value if kind in ("date", "datetime") else value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if kind == "int":
            return int(value)
        value = str(int(value)) if float(value).is_integer() else str(value)
    value = ("" if value is None else str(value)).strip()
    if value == "" or value.lower() in ("null", "none", "n/a", "-"):
        return None
    if kind == "int":
        return int(float(value.replace(",", "")))
    if kind == "date":
        return dateparser.parse(value).date()
    if kind == "datetime":
        return dateparser.parse(value).replace(tzinfo=None)
    if kind == "str":
        return value[:512]
    return value


def _read_csv(file_bytes):
    """Returns (headers, iterator of (line number, row dict))."""
    reader = csv.DictReader(io.StringIO(file_bytes.decode("utf-8-sig", errors="replace")))
    return reader.fieldnames, enumerate(reader, start=2)


def _read_excel(file_bytes):
    """Reads the first worksheet; the first non-empty row is the header row."""
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    rows = wb.worksheets[0].iter_rows(values_only=True)
    headers, header_line = None, 0
    for header_line, values in enumerate(rows, start=1):
        if any(v not in (None, "") for v in values):
            headers = [str(v).strip() if v is not None else "" for v in values]
            break

    def body():
        for line_no, values in enumerate(rows, start=header_line + 1):
            if any(v not in (None, "") for v in values):
                yield line_no, dict(zip(headers, values))
        wb.close()

    return headers, body()


def _read_rows(file_bytes, filename):
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == ".csv":
        headers, rows = _read_csv(file_bytes)
    elif ext in EXCEL_EXTENSIONS:
        try:
            headers, rows = _read_excel(file_bytes)
        except Exception as e:  # noqa: BLE001 — corrupt / password-protected workbooks
            raise ValueError(f"Could not read Excel file: {e}") from e
    elif ext == ".xls":
        raise ValueError("Old .xls files are not supported — save the file as .xlsx and upload again.")
    else:
        raise ValueError(f"Unsupported file type '{ext or filename}'. Upload a .csv or .xlsx file.")
    if not headers:
        raise ValueError("File has no header row.")

    mapping = {h: HEADER_MAP[_norm(h)] for h in headers if h and _norm(h) in HEADER_MAP}
    if not any(col == "alert_id" for col, _ in mapping.values()):
        raise ValueError("File must contain an 'Alert ID' column.")
    if not any(col == "app_id" for col, _ in mapping.values()):
        raise ValueError("File must contain an 'App ID' column.")
    return mapping, rows


# ---------------------------------------------------------------- import runs

def get_or_create_draft(session):
    run = session.scalar(select(ImportRun).where(ImportRun.status == "draft"))
    if run:
        return run
    run = ImportRun(status="draft")
    session.add(run)
    try:
        session.flush()
    except IntegrityError:  # another request created the draft at the same moment
        session.rollback()
        run = session.scalar(select(ImportRun).where(ImportRun.status == "draft"))
    return run


def stage_file(session, run, file_bytes, filename):
    """Parses a file into the run's staging rows. Live data is not touched.

    The same Alert ID in a later file of the same run replaces the earlier one.
    """
    if run.status != "draft":
        raise ValueError("This import has already been finalized or discarded.")
    mapping, rows = _read_rows(file_bytes, filename)

    log = UploadLog(filename=filename, mode="staged", run_id=run.id)
    session.add(log)
    session.flush()

    st = StagedVulnerability.__table__
    errors, batch = [], {}
    read = staged = rejected = 0

    def flush_batch():
        nonlocal staged
        if not batch:
            return
        values = list(batch.values())
        stmt = insert(st).values(values)
        session.execute(stmt.on_conflict_do_update(
            constraint="uq_staged_run_alert",
            set_={c: stmt.excluded[c] for c in values[0] if c not in ("run_id", "alert_id")},
        ))
        staged += len(values)
        batch.clear()

    all_cols = {name for _, name, _ in VULN_FIELDS}
    for line_no, raw in rows:
        read += 1
        row, row_err = {c: None for c in all_cols}, None
        for header, (col, kind) in mapping.items():
            try:
                row[col] = _convert(raw.get(header), kind)
            except (ValueError, OverflowError):
                row_err = f"line {line_no}: bad value for '{header}': {raw.get(header)!r}"
                break
        if not row_err and not row["alert_id"]:
            row_err = f"line {line_no}: missing Alert ID"
        if row_err:
            rejected += 1
            if len(errors) < 50:
                errors.append(row_err)
            continue
        row.update(run_id=run.id, upload_id=log.id, line_no=line_no)
        # Last occurrence wins when the same Alert ID appears twice in one batch.
        batch[row["alert_id"]] = row
        if len(batch) >= BATCH_SIZE:
            flush_batch()
    flush_batch()

    log.rows_read = read
    log.rows_loaded = staged
    log.rows_rejected = rejected
    log.errors = "\n".join(errors) or None
    session.commit()
    return log


# A staged row against the stored one (v), by Last Scan. NULL scan dates count as "newer"
# so data without scan dates still flows through.
_NEWER = "(s.last_scan IS NULL OR v.last_scan IS NULL OR s.last_scan > v.last_scan)"
_SAME_OR_NEWER = "(s.last_scan IS NULL OR v.last_scan IS NULL OR s.last_scan >= v.last_scan)"

# Open alerts to remediate: on an asset this run scanned more recently than the alert was last
# seen, and missing from the run. Assets the run did not include are never touched.
_REMEDIATE_FROM = """
    FROM vulnerabilities v
    JOIN (SELECT asset_id, max(last_scan) AS scan FROM staged_vulnerabilities
          WHERE run_id = :run AND asset_id IS NOT NULL GROUP BY asset_id) rs ON rs.asset_id = v.asset_id
    WHERE v.status = 'open'
      AND NOT EXISTS (SELECT 1 FROM staged_vulnerabilities s WHERE s.run_id = :run AND s.alert_id = v.alert_id)
      AND (v.last_scan IS NULL OR rs.scan IS NULL OR rs.scan > v.last_scan)
"""


def preview_run(session, run, sample_limit=200):
    """What finalizing would do, without changing anything."""
    p = {"run": run.id}
    counts = session.execute(text(f"""
        SELECT
          count(*) FILTER (WHERE v.id IS NULL) AS new,
          count(*) FILTER (WHERE v.status = 'open' AND {_SAME_OR_NEWER}) AS updated,
          count(*) FILTER (WHERE v.status = 'remediated' AND {_NEWER}) AS reopened,
          count(*) FILTER (WHERE v.id IS NOT NULL AND NOT (
              (v.status = 'open' AND {_SAME_OR_NEWER}) OR (v.status = 'remediated' AND {_NEWER}))) AS skipped,
          count(DISTINCT s.asset_id) AS assets,
          count(*) AS staged
        FROM staged_vulnerabilities s LEFT JOIN vulnerabilities v ON v.alert_id = s.alert_id
        WHERE s.run_id = :run
    """), p).one()._asdict()
    counts["remediated"] = session.scalar(text("SELECT count(*) " + _REMEDIATE_FROM), p)
    remediate_sample = session.execute(text(f"""
        SELECT v.app_id, v.asset_id, v.image, v.image_tag, v.cve_id, v.severity, v.package,
               v.last_scan AS last_seen_scan, rs.scan AS new_scan {_REMEDIATE_FROM}
        ORDER BY v.app_id, v.asset_id, v.cve_id LIMIT :lim
    """), dict(p, lim=sample_limit)).mappings().all()
    remediate_by_app = session.execute(text(f"""
        SELECT v.app_id, count(*) AS n {_REMEDIATE_FROM} GROUP BY v.app_id ORDER BY n DESC
    """), p).all()
    reopen_sample = session.execute(text(f"""
        SELECT v.app_id, v.asset_id, v.cve_id, v.severity, v.remediated_at, s.last_scan AS new_scan
        FROM staged_vulnerabilities s JOIN vulnerabilities v ON v.alert_id = s.alert_id
        WHERE s.run_id = :run AND v.status = 'remediated' AND {_NEWER}
        ORDER BY v.app_id, v.cve_id LIMIT :lim
    """), dict(p, lim=sample_limit)).mappings().all()
    return dict(counts=counts, remediate_sample=remediate_sample, remediate_by_app=remediate_by_app,
                reopen_sample=reopen_sample)


def finalize_run(session, run):
    """Applies the run: insert new alerts, update/reopen known ones, mark remediated, clear staging."""
    if run.status != "draft":
        raise ValueError("This import has already been finalized or discarded.")
    session.execute(text("SELECT pg_advisory_xact_lock(724002)"))  # one finalize at a time
    counts = preview_run(session, run, sample_limit=0)["counts"]
    now = datetime.utcnow()
    p = {"run": run.id, "now": now}
    cols = [name for _, name, _ in VULN_FIELDS]
    col_list = ", ".join(cols)

    # Remediate first: it only looks at alerts NOT in the run, so it is independent of the updates.
    session.execute(text(f"""
        UPDATE vulnerabilities SET status = 'remediated', remediated_at = :now, remediated_run_id = :run
        WHERE id IN (SELECT v.id {_REMEDIATE_FROM})
    """), p)
    session.execute(text(f"""
        UPDATE vulnerabilities v SET {", ".join(f"{c} = s.{c}" for c in cols if c != "alert_id")},
               upload_id = s.upload_id, run_id = :run, uploaded_at = :now, last_seen_at = :now,
               reopened_at = CASE WHEN v.status = 'remediated' THEN :now ELSE v.reopened_at END,
               status = 'open', remediated_at = NULL, remediated_run_id = NULL
        FROM staged_vulnerabilities s
        WHERE s.run_id = :run AND s.alert_id = v.alert_id
          AND ((v.status = 'open' AND {_SAME_OR_NEWER}) OR (v.status = 'remediated' AND {_NEWER}))
    """), p)
    session.execute(text(f"""
        INSERT INTO vulnerabilities ({col_list}, upload_id, run_id, uploaded_at, first_seen_at, last_seen_at, status)
        SELECT {", ".join("s." + c for c in cols)}, s.upload_id, :run, :now, :now, :now, 'open'
        FROM staged_vulnerabilities s
        WHERE s.run_id = :run AND NOT EXISTS (SELECT 1 FROM vulnerabilities v WHERE v.alert_id = s.alert_id)
    """), p)
    session.execute(delete(StagedVulnerability).where(StagedVulnerability.run_id == run.id))

    run.status, run.finalized_at = "finalized", now
    run.rows_new, run.rows_updated, run.rows_reopened = counts["new"], counts["updated"], counts["reopened"]
    run.rows_skipped, run.rows_remediated = counts["skipped"], counts["remediated"]
    run.assets_scanned = counts["assets"]
    session.commit()
    return run


def discard_run(session, run):
    if run.status != "draft":
        raise ValueError("Only a draft import can be discarded.")
    session.execute(delete(StagedVulnerability).where(StagedVulnerability.run_id == run.id))
    run.status = "discarded"
    session.commit()
